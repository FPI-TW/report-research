"""文字抽取門面：PDF / docx → 純文字 ＋ 抽取版本 ＋ 品質欄位，並偵測掃描檔。

docs/EXTRACTION.md §4。**簽章不變**：`extract_text(path) -> ExtractResult`，
既有四個呼叫端（`extract_all`／`sync_new_reports`／`ingest_all` 經快取／`tag_all_cli` 經快取）
一個字都不用改；新欄位全部有預設值。

## 兩條路徑，由 `EXTRACTOR` 決定（預設 `pypdf`）

- `pypdf`：現況。唯一的改動是**逐頁例外不再無聲吞掉**——失敗的頁碼記進 `pages_failed`
  （診斷 #2：「抽到 3 頁」與「抽到 30 頁」在下游長得一模一樣）。
- `pdfplumber`：`app/services/extraction/layout.py` 的版面層 → `serialize_with_index`。
  文字之外多出 Block 索引（哪一段是表格／標題，供 `extract_takeaways` 濾表格）、
  頁級失敗、品質指標（`quality.measure`，零 LLM）。

## 退回規則（E1 共識第 6 條）

pdfplumber **拋例外、整份開不起來、或抽出字數低於 `MIN_TEXT_CHARS`** 才逐檔退回 pypdf；
`extractor` 誠實記 `pypdf`、`quality_flags` 記 `fallback_from`／`fallback_reason`。
**刻意不做「兩邊字數比一比取多的」**：附錄 D 已證明原始字數比是假訊號——每頁都有
頁眉的檔案 pdfplumber 刻意丟掉頁首頁尾，字數必然較少，拿字數當判準會讓它們永遠退回。

## docx

維持 `python-docx` 段落抽取（34 篇，§3.2）。`pdfplumber` 模式下包成 Block 模型（每段一個
`paragraph`，無 bbox），`extractor` 記 `python-docx`；`pypdf` 模式下維持現況的單換行接法。
"""

from __future__ import annotations

import hashlib
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from pypdf import PdfReader

try:
    import docx  # python-docx
except ImportError:  # pragma: no cover
    docx = None

logger = logging.getLogger(__name__)

# 低於此字元數視為掃描檔/空白（無 OCR → 後續跳過嵌入）
MIN_TEXT_CHARS = 100

PYPDF_VERSION = f"pypdf-{getattr(__import__('pypdf'), '__version__', 'unknown')}"


@dataclass
class ExtractResult:
    file_hash: str
    text: str
    char_count: int
    scanned: bool
    language: str  # "zh" / "en" / "mixed"
    error: str | None = None
    # ── E1a 新增，全部有預設值，既有呼叫端不受影響 ──
    extractor: str = "pypdf"  # 'pypdf' | 'pdfplumber' | 'python-docx'
    extraction_version: str = PYPDF_VERSION
    page_count: int | None = None
    pages_failed: tuple[int, ...] = ()
    # quality.measure 的指標；pypdf 路徑只有 pages_failed 相關欄位。
    quality: dict[str, Any] = field(default_factory=dict)
    # (type, page_no, start, end)，序列化文字裡每個 Block 的區間；pypdf 路徑為空。
    blocks: tuple[tuple[str, int, int, int], ...] = ()


def file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def _detect_language(text: str) -> str:
    sample = text[:4000]
    cjk = len(re.findall(r"[一-鿿]", sample))
    latin = len(re.findall(r"[A-Za-z]", sample))
    if cjk == 0 and latin == 0:
        return "unknown"
    if cjk > latin * 2:
        return "zh"
    if latin > cjk * 2:
        return "en"
    return "mixed"


# ── pypdf 路徑 ──────────────────────────────────────────────────────────────


def _extract_pdf(path: Path) -> tuple[str, int, tuple[int, ...]]:
    """回傳 (文字, 頁數, 失敗頁碼)。失敗的頁**仍然跳過**（維持現況的輸出），但不再無聲。"""
    reader = PdfReader(str(path))
    parts: list[str] = []
    failed: list[int] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            parts.append(page.extract_text() or "")
        except Exception as exc:
            failed.append(i)
            logger.debug("pypdf page %d of %s failed: %s", i, path.name, exc)
    return "\n".join(parts), len(reader.pages), tuple(failed)


def _extract_docx(path: Path) -> str:
    if docx is None:
        return ""
    document = docx.Document(str(path))
    return "\n".join(p.text for p in document.paragraphs)


def _finish(file_hash: str, text: str, normalize: bool = True, **extra: Any) -> ExtractResult:
    # pdfplumber 路徑傳 normalize=False：Block 索引是對序列化字串算的，這裡再 strip 或
    # 折疊空行會讓區間整段位移；serialize 本來就不會產生連續三個換行或首尾空白。
    if normalize:
        text = re.sub(r"\n{3,}", "\n\n", text).strip()
    char_count = len(text)
    return ExtractResult(
        file_hash=file_hash,
        text=text,
        char_count=char_count,
        scanned=char_count < MIN_TEXT_CHARS,
        language=_detect_language(text),
        **extra,
    )


def _via_pypdf(path: Path, file_hash: str, suffix: str, flags: dict[str, Any] | None = None) -> ExtractResult:
    quality: dict[str, Any] = dict(flags or {})
    try:
        if suffix == ".pdf":
            text, n_pages, failed = _extract_pdf(path)
            if failed:
                quality["pages_failed_ratio"] = round(len(failed) / n_pages, 4) if n_pages else 1.0
            return _finish(
                file_hash, text, extractor="pypdf", extraction_version=PYPDF_VERSION,
                page_count=n_pages, pages_failed=failed, quality=quality,
            )
        if suffix in (".docx", ".doc"):
            return _finish(
                file_hash, _extract_docx(path), extractor="python-docx", extraction_version="python-docx",
                quality=quality,
            )
        return _finish(file_hash, "", quality=quality)  # 圖片等：視為掃描檔
    except Exception as exc:  # 損毀檔不要中斷整批
        return ExtractResult(file_hash, "", 0, True, "unknown", error=str(exc), quality=quality)


# ── pdfplumber 路徑 ─────────────────────────────────────────────────────────


def _docx_document(path: Path):
    """docx → Document：每段一個 paragraph Block，無 bbox（§3.2）。"""
    from app.services.extraction import EXTRACTION_VERSION
    from app.services.extraction.model import Block, Document, Page

    if docx is None:
        return Document(extractor="python-docx", extraction_version=EXTRACTION_VERSION, error="python-docx 未安裝")
    document = docx.Document(str(path))
    blocks = tuple(
        Block(type="paragraph", page_no=1, order=i, bbox=None, text=p.text)
        for i, p in enumerate(document.paragraphs)
        if p.text.strip()
    )
    return Document(
        extractor="python-docx",
        extraction_version=EXTRACTION_VERSION,
        pages=(Page(1, 0.0, 0.0, blocks),),
    )


def _measure_with_coverage(doc, path: Path):
    """`quality.measure` 加上 `layout_coverage`：需要 pypdfium2 逐頁 render。

    v3 生產路徑從未傳 renderer，coverage 全庫是 null。實測 32 頁只多 0.4 秒（抽取本身
    4.9 秒），而它是唯一對「整欄漏抽」敏感的指標。render 開不起來就退回不算 coverage
    （fail-open）：品質量測不可以擋住抽取。"""
    from app.services.extraction.quality import measure

    try:
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(str(path))
    except Exception as exc:
        logger.warning("pypdfium2 open failed on %s; layout_coverage skipped: %s", path.name, exc)
        return measure(doc)
    try:
        return measure(doc, renderer=pdf)
    finally:
        try:
            pdf.close()
        except Exception:  # pragma: no cover — 版本差異
            pass


def _via_pdfplumber(path: Path, file_hash: str, suffix: str) -> ExtractResult:
    from app.services.extraction.model import serialize_with_index
    from app.services.extraction.quality import measure

    if suffix not in (".pdf", ".docx", ".doc"):
        return _finish(file_hash, "")
    try:
        if suffix == ".pdf":
            from app.services.extraction.layout import extract_document

            doc = extract_document(path)
        else:
            doc = _docx_document(path)
        if doc.error:
            raise RuntimeError(doc.error)
        text, spans = serialize_with_index(doc)
        q = _measure_with_coverage(doc, path) if suffix == ".pdf" else measure(doc)
    except Exception as exc:
        # 退回 pypdf。原因記進 quality_flags，extractor 誠實記 pypdf。
        logger.warning("pdfplumber failed on %s, falling back to pypdf: %s", path.name, exc)
        return _via_pypdf(
            path, file_hash, suffix, flags={"fallback_from": "pdfplumber", "fallback_reason": f"error: {exc}"[:200]},
        )

    res = _finish(
        file_hash, text, normalize=False, extractor=doc.extractor, extraction_version=doc.extraction_version,
        page_count=doc.page_count, pages_failed=doc.pages_failed, quality=q.as_flags(),
        blocks=tuple((s.type, s.page_no, s.start, s.end) for s in spans),
    )
    if res.scanned:
        # 抽不出字才退回；字數多寡不比（見檔頭）。
        fallback = _via_pypdf(
            path, file_hash, suffix, flags={"fallback_from": doc.extractor, "fallback_reason": "below_min_chars"},
        )
        if not fallback.scanned:
            return fallback
        # 兩邊都抽不到：維持 pdfplumber 的結果（它的 pages_failed／品質欄位比較完整）
        res.quality["fallback_from"] = doc.extractor
        res.quality["fallback_reason"] = "below_min_chars_both"
    return res


# ── 門面 ────────────────────────────────────────────────────────────────────


def extract_text(path: Path, extractor: str | None = None) -> ExtractResult:
    """抽取入口。`extractor` 省略時讀設定的 `EXTRACTOR`（預設 pypdf）。"""
    file_hash = file_sha256(path)
    suffix = path.suffix.lower()
    if extractor is None:
        from app.config import get_settings

        extractor = get_settings().extractor
    if extractor == "pdfplumber":
        return _via_pdfplumber(path, file_hash, suffix)
    return _via_pypdf(path, file_hash, suffix)
